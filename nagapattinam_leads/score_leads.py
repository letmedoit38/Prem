"""
Task 4 — Lead Scoring & Master Database Builder
Merges MCA, Udyam, and GeM data; applies scoring rules; assigns IOB products.
Saves master_leads.csv and master_leads.xlsx (with conditional formatting).
"""

import logging
import os
from datetime import datetime

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

CUTOFF_DATE = datetime(2024, 1, 1)


# ── Scoring ───────────────────────────────────────────────────────────────────

def _parse_date(val) -> datetime | None:
    if not val or str(val).strip() in ("", "nan", "NaT", "None"):
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d-%b-%Y",
                "%Y/%m/%d", "%b %d, %Y", "%d %b %Y"):
        try:
            return datetime.strptime(str(val).strip(), fmt)
        except ValueError:
            pass
    return None


def _score_row(row: pd.Series) -> int:
    score = 0
    rules = config.SCORING_RULES

    # Authorized capital > ₹10 Lakh (1,000,000)
    try:
        cap = float(str(row.get("AUTHORIZED_CAPITAL", 0)).replace(",", "") or 0)
        if cap > 1_000_000:
            score += rules["authorized_capital_above_10L"]
    except (ValueError, TypeError):
        pass

    # Incorporated after Jan 2024
    inc_date = _parse_date(row.get("INCORPORATION_DATE", ""))
    if inc_date and inc_date >= CUTOFF_DATE:
        score += rules["incorporated_after_jan_2024"]

    # Has GST
    gst_val = str(row.get("HAS_GST", "")).strip().lower()
    if gst_val in ("true", "yes", "1", "y"):
        score += rules["has_gst"]

    # Sector bonus
    sector = str(row.get("SECTOR", "")).strip()
    if "manufacturing" in sector.lower():
        score += rules["sector_manufacturing"]
    elif sector.lower() in ("services", "technology", "it"):
        score += rules["sector_services"]

    # GeM registered
    gem_val = row.get("GEM_REGISTERED", False)
    if gem_val is True or str(gem_val).lower() in ("true", "yes", "1"):
        score += rules["gem_registered"]

    # Has contact info
    has_email = str(row.get("EMAIL", "")).strip() not in ("", "nan", "None")
    has_phone = str(row.get("PHONE", "")).strip() not in ("", "nan", "None")
    if has_email or has_phone:
        score += rules["has_contact_info"]

    # Normalise to 0–100
    raw_max = config.SCORE_MAX
    return round(min(score / raw_max * 100, 100))


def _assign_product(row: pd.Series) -> str:
    sector = str(row.get("SECTOR", "")).strip().title()
    for key in config.BANK_PRODUCTS:
        if key.lower() in sector.lower():
            return config.BANK_PRODUCTS[key]
    return config.BANK_PRODUCTS["Other"]


# ── Excel export with formatting ──────────────────────────────────────────────

def _export_xlsx(df: pd.DataFrame, path: str) -> None:
    df.to_excel(path, index=False, engine="openpyxl")

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.title = "Master Leads"

    # Header row styling
    header_fill = PatternFill("solid", fgColor="003366")   # IOB dark blue
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Auto-size columns
    for col_idx, col in enumerate(df.columns, start=1):
        max_len = max(
            len(str(col)),
            df[col].astype(str).str.len().max() if len(df) > 0 else 0,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    # Colour-code SCORE column
    score_col_idx = list(df.columns).index("SCORE") + 1
    score_col_letter = get_column_letter(score_col_idx)

    for row_idx in range(2, ws.max_row + 1):
        cell = ws[f"{score_col_letter}{row_idx}"]
        try:
            val = int(cell.value or 0)
        except (ValueError, TypeError):
            val = 0

        if val >= 75:
            cell.fill = PatternFill("solid", fgColor="00B050")   # green
            cell.font = Font(bold=True, color="FFFFFF")
        elif val >= 50:
            cell.fill = PatternFill("solid", fgColor="FFC000")   # amber
        else:
            cell.fill = PatternFill("solid", fgColor="FF0000")   # red
            cell.font = Font(color="FFFFFF")

    # Freeze top row
    ws.freeze_panes = "A2"

    # Thin borders on all cells
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


# ── Public Entry Point ────────────────────────────────────────────────────────

def build_master_leads(
    df_mca: pd.DataFrame,
    df_mfg: pd.DataFrame,
    df_svc: pd.DataFrame,
    df_gem: pd.DataFrame,
) -> pd.DataFrame:
    """
    Merge all source DataFrames, score each lead, assign IOB products,
    sort descending by score. Saves master_leads.csv and master_leads.xlsx.
    """
    os.makedirs(config.DATA_DIR, exist_ok=True)

    frames = []
    for label, df in [("MCA", df_mca), ("MFG", df_mfg),
                       ("SVC", df_svc), ("GeM", df_gem)]:
        if df is not None and not df.empty:
            frames.append(df)
            print(f"  [{label}] {len(df)} records to merge")
        else:
            print(f"  [{label}] 0 records (skipped)")

    if not frames:
        print("\n[Score] WARNING: No data from any source. "
              "Run fetchers first or check API keys.")
        return pd.DataFrame(columns=config.MASTER_COLUMNS)

    master = pd.concat(frames, ignore_index=True)

    # Ensure all required columns exist
    for col in config.MASTER_COLUMNS:
        if col not in master.columns:
            master[col] = ""

    # Deduplicate across sources: prefer MCA CIN, then Udyam No
    master = master[master["COMPANY_NAME"].str.strip() != ""]
    master = master.drop_duplicates(subset=["CIN_OR_UDYAM_NO", "DISTRICT"],
                                    keep="first")

    # Score
    print(f"\n[Score] Scoring {len(master)} leads …")
    master["SCORE"] = master.apply(_score_row, axis=1)

    # Assign products
    master["BANK_PRODUCT"] = master.apply(_assign_product, axis=1)

    # Sort
    master = master.sort_values("SCORE", ascending=False).reset_index(drop=True)

    # Reorder to canonical column order
    extra_cols = [c for c in master.columns if c not in config.MASTER_COLUMNS]
    master = master[config.MASTER_COLUMNS + extra_cols]

    # Save CSV
    master.to_csv(config.MASTER_LEADS_CSV_OUT, index=False, encoding="utf-8-sig")
    print(f"[Score] Saved CSV  → {config.MASTER_LEADS_CSV_OUT}")

    # Save XLSX with formatting
    _export_xlsx(master, config.MASTER_LEADS_XLSX_OUT)
    print(f"[Score] Saved XLSX → {config.MASTER_LEADS_XLSX_OUT}")

    # Summary
    print(f"\n{'='*55}")
    print(f"  MASTER LEADS SUMMARY")
    print(f"{'='*55}")
    print(f"  Total leads     : {len(master)}")
    print(f"  High priority   : {(master['SCORE'] >= 75).sum()}  (score ≥ 75)")
    print(f"  Medium priority : {((master['SCORE'] >= 50) & (master['SCORE'] < 75)).sum()}  (score 50–74)")
    print(f"  Low priority    : {(master['SCORE'] < 50).sum()}  (score < 50)")
    print(f"\n  By district:")
    for dist in config.DISTRICTS:
        cnt = (master["DISTRICT"] == dist).sum()
        print(f"    {dist:<20} {cnt}")
    print(f"\n  By sector:")
    for sec, grp in master.groupby("SECTOR"):
        print(f"    {sec:<20} {len(grp)}")
    print(f"\n  TOP 5 LEADS:")
    print(f"{'='*55}")
    top5 = master.head(5)
    for _, r in top5.iterrows():
        print(f"  {r['SCORE']:>3}/100 | {r['COMPANY_NAME'][:35]:<35} | "
              f"{r['DISTRICT']:<16} | {r['SECTOR']}")
    print(f"{'='*55}")

    return master


# ── CLI runner ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)

    # Load previously saved CSVs if they exist
    def _load(path):
        return pd.read_csv(path) if os.path.exists(path) else pd.DataFrame()

    df_mca = _load(config.MCA_CSV_OUT)
    df_mfg = _load(config.UDYAM_MFG_CSV_OUT)
    df_svc = _load(config.UDYAM_SVC_CSV_OUT)
    df_gem = _load(config.GEM_CSV_OUT)

    master = build_master_leads(df_mca, df_mfg, df_svc, df_gem)
    print(f"\nMaster leads: {len(master)} rows saved.")
